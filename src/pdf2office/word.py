"""Editable Word output from the PDF intermediate representation.

``layout`` writes positioned Word text boxes and VML paths, not page screenshots.
``flow`` writes ordinary paragraphs and separate inline vector/image canvases.
VML is supported by desktop Microsoft Word; other DOCX readers may render it
differently. Neither mode reconstructs semantic CAD objects or Word tables.
"""

from __future__ import annotations

import io
import math
import re
from pathlib import Path
from statistics import median
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw

if TYPE_CHECKING:
    from .fonts import FontCatalog
    from .models import Drawing, PdfPage, TextSpan

_VML = "urn:schemas-microsoft-com:vml"
_OFFICE = "urn:schemas-microsoft-com:office:office"
_WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MAX_WORD_POINTS = 22 * 72
_XML_UNSAFE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _v(tag: str, **attributes: str) -> etree._Element:
    return etree.Element(f"{{{_VML}}}{tag}", attributes, nsmap={"v": _VML, "o": _OFFICE})


def _clean(text: str) -> str:
    return _XML_UNSAFE.sub("\ufffd", text)


def _color(color: tuple[int, int, int]) -> str:
    return "#" + "".join(f"{max(0, min(255, int(c))):02X}" for c in color)


def _position(x: float, y: float, width: float, height: float, z: int) -> str:
    return (
        f"position:absolute;margin-left:{x:.4f}pt;margin-top:{y:.4f}pt;"
        f"width:{max(width, 0.01):.4f}pt;height:{max(height, 0.01):.4f}pt;"
        f"z-index:{z};mso-position-horizontal-relative:page;"
        "mso-position-vertical-relative:page;mso-wrap-style:none"
    )


def _pict(paragraph, shape: etree._Element) -> None:
    pict = OxmlElement("w:pict")
    pict.append(shape)
    paragraph.add_run()._r.append(pict)


def _drawing_bounds(drawing: Drawing) -> tuple[float, float, float, float] | None:
    points = [p for segment in drawing.segments for p in segment.points]
    if not points:
        return None
    xs, ys = zip(*points)
    return min(xs), min(ys), max(xs), max(ys)


def _drawing_shape(
    drawing: Drawing,
    shape_id: int,
    scale: float,
    *,
    group_origin: tuple[float, float] | None = None,
) -> etree._Element | None:
    bounds = _drawing_bounds(drawing)
    if bounds is None:
        return None
    left, top, right, bottom = bounds
    width, height = max(right - left, 0.01), max(bottom - top, 0.01)
    factor = 1000

    def point(p: tuple[float, float]) -> str:
        return f"{round((p[0] - left) * factor)},{round((p[1] - top) * factor)}"

    parts: list[str] = []
    last = None
    for segment in drawing.segments:
        if segment.kind not in {"line", "cubic"}:
            raise ValueError(f"Unsupported drawing segment: {segment.kind}")
        expected = 2 if segment.kind == "line" else 4
        if len(segment.points) != expected:
            raise ValueError(f"{segment.kind} requires {expected} control points")
        start = segment.points[0]
        if last != start:
            if last is not None and drawing.closed:
                parts.append("x")
            parts.append(f"m {point(start)}")
        if segment.kind == "line":
            parts.append(f"l {point(segment.points[1])}")
        else:
            parts.append("c " + ",".join(point(p) for p in segment.points[1:]))
        last = segment.points[-1]
    if drawing.closed:
        parts.append("x")
    parts.append("e")
    if group_origin is None:
        style = _position(left * scale, top * scale, width * scale, height * scale, shape_id)
    else:
        style = (
            f"position:absolute;left:{round((left - group_origin[0]) * factor)};"
            f"top:{round((top - group_origin[1]) * factor)};"
            f"width:{round(width * factor)};height:{round(height * factor)}"
        )
    shape = _v(
        "shape",
        id=f"PDFShape{shape_id}",
        style=style,
        coordorigin="0,0",
        coordsize=f"{round(width * factor)},{round(height * factor)}",
        path=" ".join(parts),
        stroked="t" if drawing.stroke is not None else "f",
        filled="t" if drawing.fill is not None else "f",
        strokeweight=f"{max(drawing.width * scale, 0.1):.4f}pt",
    )
    if drawing.stroke is not None:
        shape.set("strokecolor", _color(drawing.stroke))
        dashes = getattr(drawing, "dashes", ())
        if dashes:
            # PDF measures dashes in page points; VML measures each dash/gap
            # in multiples of the rendered stroke width. An odd PDF pattern
            # repeats twice before its dash/gap roles repeat.
            pattern = tuple(dashes) * (2 if len(dashes) % 2 else 1)
            stroke_width = max(drawing.width * scale, 0.1)
            dashstyle = " ".join(f"{length * scale / stroke_width:.6g}" for length in pattern)
            shape.append(_v("stroke", dashstyle=dashstyle))
    if drawing.fill is not None:
        shape.set("fillcolor", _color(drawing.fill))
    return shape


class _Fonts:
    def __init__(self, overrides: dict[str, str] | None, warnings: list[str], catalog: FontCatalog | None = None):
        from .fonts import FontCatalog

        self.catalog = catalog if catalog is not None else FontCatalog()
        self.overrides = overrides or {}
        self.warnings = warnings
        self.cache: dict[tuple[str, bool, bool], str] = {}

    def family(self, span: TextSpan) -> str:
        original = span.font or "Arial"
        requested = self.overrides.get(
            original,
            self.overrides.get(original.split("+", 1)[-1], self.overrides.get("*", original)),
        )
        key = (requested, bool(span.bold), bool(span.italic))
        if key not in self.cache:
            match = self.catalog.resolve(requested, bold=span.bold, italic=span.italic)
            self.cache[key] = match.family
            if match.substituted:
                self.warnings.append(f"DOCX: шрифт «{requested}» заменён на «{match.family}».")
        return self.cache[key]


def _format_run(run, span: TextSpan, fonts: _Fonts, scale: float = 1.0) -> None:
    family = fonts.family(span)
    run.font.name = family
    run.font.size = Pt(max(1, min(1638, span.size * scale)))
    run.bold = bool(span.bold)
    run.italic = bool(span.italic)
    run.font.color.rgb = RGBColor(*(max(0, min(255, int(c))) for c in span.color))
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), family)


def _textbox(span: TextSpan, shape_id: int, scale: float, fonts: _Fonts) -> etree._Element:
    from docx.text.paragraph import Paragraph

    x0, y0, x1, y1 = span.bbox
    width, height = x1 - x0, y1 - y0
    # PDF boxes enclose the already rotated glyphs; VML rotates an unrotated
    # box about its centre. Undo quarter-turn dimensions before applying it.
    if math.isclose(abs(span.angle) % 180, 90, abs_tol=0.01):
        centre_x, centre_y = (x0 + x1) / 2, (y0 + y1) / 2
        width, height = height, width
        x0, y0 = centre_x - width / 2, centre_y - height / 2
    style = _position(x0 * scale, y0 * scale, width * scale, height * scale, shape_id)
    if span.angle:
        style += f";rotation:{span.angle:.4f}"
    shape = _v("rect", id=f"PDFText{shape_id}", style=style, filled="f", stroked="f")
    textbox = _v("textbox", inset="0,0,0,0", style="mso-fit-shape-to-text:t")
    content = OxmlElement("w:txbxContent")
    element = OxmlElement("w:p")
    content.append(element)
    paragraph = Paragraph(element, None)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(max(span.size * scale * 1.1, 1))
    _format_run(paragraph.add_run(_clean(span.text)), span, fonts, scale)
    textbox.append(content)
    shape.append(textbox)
    return shape


def _clean_ocr_raster(raster, spans: list[TextSpan]) -> tuple[bytes, bool]:
    """Clear OCR text boxes in a copy using their surrounding median colour.

    This deliberately simple replacement cannot reconstruct geometry hidden by
    letters. It only affects exported image pixels and does not edit PDF input.
    """
    x0, y0, x1, y1 = raster.bbox
    if x1 <= x0 or y1 <= y0:
        return raster.data, False
    boxes = [
        span.bbox for span in spans
        if span.from_ocr and span.text.strip()
        and span.bbox[0] < x1 and span.bbox[2] > x0
        and span.bbox[1] < y1 and span.bbox[3] > y0
    ]
    if not boxes:
        return raster.data, False
    with Image.open(io.BytesIO(raster.data)) as opened:
        source = opened.convert("RGBA" if "A" in opened.getbands() or "transparency" in opened.info else "RGB")
    cleaned = source.copy()
    draw = ImageDraw.Draw(cleaned)
    pixels = source.load()
    width, height = source.size
    scale_x, scale_y = width / (x1 - x0), height / (y1 - y0)
    changed = False
    for box in boxes:
        left = max(0, math.floor((box[0] - x0) * scale_x))
        top = max(0, math.floor((box[1] - y0) * scale_y))
        right = min(width, math.ceil((box[2] - x0) * scale_x))
        bottom = min(height, math.ceil((box[3] - y0) * scale_y))
        if right <= left or bottom <= top:
            continue
        samples = []
        # Sample only outside the replaced box, always from the original image,
        # so neighbouring OCR boxes cannot contaminate later replacements.
        for y in (top - 2, top - 1, bottom, bottom + 1):
            if 0 <= y < height:
                samples.extend(pixels[x, y] for x in range(max(0, left - 2), min(width, right + 2), max(1, (right - left) // 128)))
        for x in (left - 2, left - 1, right, right + 1):
            if 0 <= x < width:
                samples.extend(pixels[x, y] for y in range(top, bottom, max(1, (bottom - top) // 64)))
        if samples:
            background = tuple(round(median(channel)) for channel in zip(*samples))
        else:
            background = (255, 255, 255, 255) if source.mode == "RGBA" else (255, 255, 255)
        draw.rectangle((left, top, right - 1, bottom - 1), fill=background)
        changed = True
    if not changed:
        return raster.data, False
    result = io.BytesIO()
    cleaned.save(result, format="PNG")
    return result.getvalue(), True


def _floating_image(paragraph, raster, scale: float, data: bytes | None = None) -> None:
    x0, y0, x1, y1 = raster.bbox
    inline = paragraph.add_run().add_picture(
        io.BytesIO(raster.data if data is None else data), width=Pt(max((x1 - x0) * scale, 0.01)),
        height=Pt(max((y1 - y0) * scale, 0.01)),
    )._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, coordinate in (("H", x0), ("V", y0)):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = str(round(coordinate * scale * 12700))
        position.append(offset)
        anchor.append(position)
    anchor.append(inline.find(qn("wp:extent")))
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)


def _flow_lines(spans: list[TextSpan]) -> list[list[TextSpan]]:
    """Group text runs by approximate baseline, preserving left-to-right order."""
    lines: list[list[TextSpan]] = []
    for span in sorted(spans, key=lambda s: (s.bbox[1], s.bbox[0])):
        baseline = span.origin[1] if span.origin else span.bbox[3]
        for line in reversed(lines):
            first = line[0]
            other = first.origin[1] if first.origin else first.bbox[3]
            if abs(baseline - other) <= max(2.0, min(span.size, first.size) * 0.25):
                line.append(span)
                break
        else:
            lines.append([span])
    for line in lines:
        line.sort(key=lambda span: span.bbox[0])
    return lines


def _flow_canvas(document, page: PdfPage, available_width: float, available_height: float, next_id: int) -> int:
    bounds = [bound for drawing in page.drawings if (bound := _drawing_bounds(drawing))]
    if not bounds:
        return next_id
    left = min(b[0] for b in bounds)
    top = min(b[1] for b in bounds)
    width = max(max(b[2] for b in bounds) - left, 1)
    height = max(max(b[3] for b in bounds) - top, 1)
    scale = min(1, available_width / width, available_height / height)
    group = _v(
        "group", id=f"PDFCanvas{next_id}",
        style=f"width:{width * scale:.4f}pt;height:{height * scale:.4f}pt",
        coordorigin="0,0", coordsize=f"{round(width * 1000)},{round(height * 1000)}",
    )
    next_id += 1
    for drawing in page.drawings:
        shape = _drawing_shape(drawing, next_id, scale, group_origin=(left, top))
        if shape is not None:
            group.append(shape)
            next_id += 1
    _pict(document.add_paragraph(), group)
    return next_id


def export_docx(
    pages: list[PdfPage],
    output: Path,
    mode: str = "layout",
    font_overrides: dict[str, str] | None = None,
    *,
    font_catalog: FontCatalog | None = None,
) -> list[str]:
    """Write an editable DOCX and return user-facing Russian limitations.

    Fonts are referenced by family name, not embedded. ``font_catalog`` can
    include user-supplied font folders. ``font_overrides`` maps
    source PDF font names to installed families; ``*`` sets a global replacement.
    Raises ``ValueError`` for an invalid mode, no pages or invalid page geometry.
    """
    if mode not in {"layout", "flow"}:
        raise ValueError("DOCX mode must be 'layout' or 'flow'")
    if not pages:
        raise ValueError("Cannot create DOCX without pages")
    for page in pages:
        if not all(math.isfinite(v) and v > 0 for v in (page.width, page.height)):
            raise ValueError(f"Invalid PDF page size on page {page.number}")

    warnings = [
        "DOCX: шрифты указаны по именам и не встроены; установите их на компьютере получателя.",
        "DOCX: сложные заливки, обрезка, прозрачность и порядок наложения объектов могут отличаться от PDF.",
    ]
    if mode == "layout":
        warnings.append(
            "DOCX: текст сохранён в редактируемых текстовых полях, линии и кривые — в фигурах VML. "
            "Рекомендуется настольный Microsoft Word; совместимость других редакторов требует проверки."
        )
    else:
        warnings.append(
            "DOCX: режим потока создаёт абзацы; многоколоночная структура и таблицы не восстанавливаются. "
            "Рисунки и изображения вынесены после текста каждой исходной страницы."
        )
    if any(page.images for page in pages):
        warnings.append("DOCX: исходные растровые изображения остаются изображениями; их линии не становятся фигурами Word.")
    if any(
        drawing.stroke is not None and getattr(drawing, "dashes", ())
        and not math.isclose(getattr(drawing, "dash_offset", 0.0), 0, abs_tol=1e-8)
        for page in pages for drawing in page.drawings
    ):
        warnings.append("DOCX: штриховые линии сохранены, но начальная фаза штрихов в VML не поддерживается и сброшена.")
    if any(page.ocr_used for page in pages):
        warnings.append("DOCX: распознанный OCR текст нужно проверить; нераспознанные символы остаются частью скана.")
        if mode == "flow":
            warnings.append("DOCX: исходный скан оставлен после редактируемого текста как отдельное изображение для сверки.")
    if any(_clean(span.text) != span.text for page in pages for span in page.texts):
        warnings.append("DOCX: неподдерживаемые XML управляющие символы заменены знаком �.")
    if mode == "layout" and any(not math.isclose(span.angle % 90, 0, abs_tol=0.01) for page in pages for span in page.texts):
        warnings.append("DOCX: положение текста с поворотом, отличным от 90°, приблизительное.")

    document = Document()
    document.core_properties.title = "PDF → Word"
    document.core_properties.subject = "Редактируемый результат преобразования PDF"
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    fonts = _Fonts(font_overrides, warnings, font_catalog)
    next_id = 1025
    for index, page in enumerate(pages):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        scale = min(1.0, _MAX_WORD_POINTS / max(page.width, page.height))
        if scale < 1:
            warnings.append(f"DOCX: страница {page.number} уменьшена до {scale:.1%}: Word ограничивает размер листа 22 дюймами.")
        section.page_width = Pt(page.width * scale)
        section.page_height = Pt(page.height * scale)
        margin = 0 if mode == "layout" else min(36, min(page.width, page.height) * scale / 10)
        section.top_margin = section.bottom_margin = Pt(margin)
        section.left_margin = section.right_margin = Pt(margin)
        section.header_distance = section.footer_distance = Pt(0)
        if mode == "layout":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(1)
            for raster in page.images:
                data, cleaned = _clean_ocr_raster(raster, page.texts)
                _floating_image(paragraph, raster, scale, data)
                if cleaned:
                    warnings.append(
                        "DOCX: под распознанным текстом растр очищен цветом окружающего фона, "
                        "чтобы текст можно было заменить. Пересекающиеся линии и сложный фон "
                        "могут быть повреждены и требуют проверки; исходный PDF не изменён."
                    )
            for drawing in page.drawings:
                shape = _drawing_shape(drawing, next_id, scale)
                if shape is not None:
                    _pict(paragraph, shape)
                    next_id += 1
            for span in page.texts:
                if span.text:
                    _pict(paragraph, _textbox(span, next_id, scale, fonts))
                    next_id += 1
        else:
            for line in _flow_lines(page.texts):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(2)
                previous = None
                for span in line:
                    gap = span.bbox[0] - previous.bbox[2] if previous else 0
                    if previous and gap > max(1, min(previous.size, span.size) * 0.15) and not previous.text.endswith(" ") and not span.text.startswith(" "):
                        paragraph.add_run(" ")
                    _format_run(paragraph.add_run(_clean(span.text)), span, fonts, scale)
                    previous = span
            available_width = page.width * scale - margin * 2
            available_height = max(1, page.height * scale - margin * 2 - 12)
            next_id = _flow_canvas(document, page, available_width, available_height, next_id)
            for raster in page.images:
                width = max(raster.bbox[2] - raster.bbox[0], 0.01)
                height = max(raster.bbox[3] - raster.bbox[1], 0.01)
                image_scale = min(1, available_width / width, available_height / height)
                document.add_paragraph().add_run().add_picture(
                    io.BytesIO(raster.data), width=Pt(width * image_scale), height=Pt(height * image_scale),
                )
            if not page.texts and not page.drawings and not page.images:
                document.add_paragraph()
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return list(dict.fromkeys(warnings))
