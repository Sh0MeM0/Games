from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw

from pdf2office.models import Drawing, PdfPage, RasterImage, Segment, TextSpan
from pdf2office.fonts import FontMatch
from pdf2office.word import export_docx


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def read_xml(path: Path):
    with ZipFile(path) as package:
        return etree.fromstring(package.read("word/document.xml"))


def sample_page() -> PdfPage:
    return PdfPage(
        number=1, width=595, height=842,
        texts=[TextSpan("Чертёж <А> & размер 25 мм", (30, 40, 260, 56), "Liberation Sans", 12, (25, 50, 75), bold=True)],
        drawings=[
            Drawing([Segment("line", ((20, 90), (220, 90)))], stroke=(255, 0, 0), width=2),
            Drawing([Segment("cubic", ((30, 110), (40, 160), (90, 160), (100, 110)))], stroke=(0, 0, 255), fill=(230, 240, 250), closed=True),
        ],
        images=[],
    )


def test_layout_has_editable_unicode_text_and_vector_paths(tmp_path):
    output = tmp_path / "layout.docx"
    warnings = export_docx([sample_page()], output)
    xml = read_xml(output)
    assert xml.xpath("//w:txbxContent//w:t/text()", namespaces=NS) == ["Чертёж <А> & размер 25 мм"]
    assert len(xml.xpath("//v:shape", namespaces=NS)) == 2
    assert "l " in xml.xpath("//v:shape/@path", namespaces=NS)[0]
    assert "c " in xml.xpath("//v:shape/@path", namespaces=NS)[1]
    assert " x e" in xml.xpath("//v:shape/@path", namespaces=NS)[1]
    assert xml.xpath("//v:shape/@strokecolor", namespaces=NS) == ["#FF0000", "#0000FF"]
    assert xml.xpath("//v:shape/@fillcolor", namespaces=NS) == ["#E6F0FA"]
    assert xml.xpath("//w:txbxContent//w:b", namespaces=NS)
    assert xml.xpath("//w:rFonts/@w:hAnsi", namespaces=NS) == ["Liberation Sans"]
    assert Document(output).sections[0].page_width.pt == pytest.approx(595)
    assert any("VML" in warning for warning in warnings)


def test_page_sizes_and_shape_positions_remain_independent(tmp_path):
    first = sample_page()
    second = PdfPage(2, 842, 595, [TextSpan("Лист 2", (40, 60, 100, 80), "Liberation Sans", 14, (0, 0, 0))], [], [])
    output = tmp_path / "sections.docx"
    export_docx([first, second], output)
    reopened = Document(output)
    assert [(s.page_width.pt, s.page_height.pt) for s in reopened.sections] == [(595, 842), (842, 595)]
    xml = read_xml(output)
    assert len(xml.xpath("//w:sectPr", namespaces=NS)) == 2
    assert "margin-top:60.0000pt" in xml.xpath("//v:rect/@style", namespaces=NS)[1]
    ids = xml.xpath("//v:shape/@id | //v:rect/@id", namespaces=NS)
    assert len(ids) == len(set(ids))


def test_flow_text_is_native_and_vectors_are_separate_editable_shapes(tmp_path):
    page = sample_page()
    page.texts = [
        TextSpan("Первая", (20, 20, 65, 34), "Liberation Sans", 12, (0, 0, 0), origin=(20, 32)),
        TextSpan("строка", (70, 20, 115, 34), "Liberation Sans", 12, (0, 0, 0), italic=True, origin=(70, 32)),
        TextSpan("Вторая строка", (20, 40, 120, 54), "Liberation Sans", 12, (0, 0, 0), origin=(20, 52)),
    ]
    output = tmp_path / "flow.docx"
    warnings = export_docx([page], output, mode="flow")
    paragraphs = Document(output).paragraphs
    assert paragraphs[0].text == "Первая строка"
    assert paragraphs[1].text == "Вторая строка"
    xml = read_xml(output)
    assert not xml.xpath("//w:txbxContent", namespaces=NS)
    assert len(xml.xpath("//v:group/v:shape", namespaces=NS)) == 2
    assert any("таблицы не восстанавливаются" in warning for warning in warnings)


@pytest.mark.parametrize("mode", ["layout", "flow"])
def test_dashed_vector_is_editable_and_uses_relative_stroke_units(tmp_path, mode):
    page = sample_page()
    drawing = page.drawings[0]
    drawing.width = 1.5
    drawing.dashes = (6.0, 3.0)
    drawing.dash_offset = 2.0
    output = tmp_path / f"dashes-{mode}.docx"
    warnings = export_docx([page], output, mode=mode)
    xml = read_xml(output)
    assert xml.xpath("//v:shape/v:stroke/@dashstyle", namespaces=NS) == ["4 2"]
    assert len(xml.xpath("//v:shape", namespaces=NS)) == 2
    assert "l " in xml.xpath("//v:shape/@path", namespaces=NS)[0]
    assert any("начальная фаза" in warning for warning in warnings)
    assert Document(output).sections[0].page_width.pt == pytest.approx(595)


def test_odd_pdf_dash_pattern_repeats_for_vml_pairs(tmp_path):
    page = sample_page()
    page.drawings[0].dashes = (4.0,)
    page.drawings[0].dash_offset = 0.0
    output = tmp_path / "odd-dash.docx"
    warnings = export_docx([page], output)
    assert read_xml(output).xpath("//v:shape/v:stroke/@dashstyle", namespaces=NS) == ["2 2"]
    assert not any("начальная фаза" in warning for warning in warnings)


def test_image_is_embedded_with_page_coordinates(tmp_path):
    data = io.BytesIO()
    Image.new("RGB", (12, 8), "red").save(data, format="PNG")
    page = PdfPage(1, 300, 400, [], [], [RasterImage((20, 30, 140, 110), data.getvalue())], ocr_used=True)
    output = tmp_path / "image.docx"
    warnings = export_docx([page], output)
    xml = read_xml(output)
    assert xml.xpath("//wp:anchor/wp:positionH/wp:posOffset/text()", namespaces=NS) == [str(20 * 12700)]
    assert xml.xpath("//wp:anchor/wp:positionV/wp:posOffset/text()", namespaces=NS) == [str(30 * 12700)]
    with ZipFile(output) as package:
        assert len([name for name in package.namelist() if name.startswith("word/media/")]) == 1
    assert any("распознанный OCR" in warning for warning in warnings)
    assert any("остаются изображениями" in warning for warning in warnings)
    assert len(Document(output).sections) == 1


def test_ocr_clears_only_recognized_text_in_exported_image(tmp_path):
    image = Image.new("RGB", (100, 100), (238, 230, 220))
    draw = ImageDraw.Draw(image)
    draw.rectangle((32, 32, 48, 48), fill=(0, 0, 0))
    draw.line((80, 10, 80, 90), fill=(0, 0, 255), width=2)
    data = io.BytesIO()
    image.save(data, format="PNG")
    original = data.getvalue()
    raster = RasterImage((10, 20, 210, 220), original)
    page = PdfPage(
        1, 300, 400,
        [TextSpan("Исправить", (70, 80, 110, 120), "Liberation Sans", 12, (0, 0, 0), from_ocr=True)],
        [], [raster], ocr_used=True,
    )
    output = tmp_path / "editable-scan.docx"
    warnings = export_docx([page], output)
    with ZipFile(output) as package:
        media = next(name for name in package.namelist() if name.startswith("word/media/"))
        cleaned = Image.open(io.BytesIO(package.read(media))).convert("RGB")
        assert cleaned.getpixel((40, 40)) == (238, 230, 220)
        for point in ((10, 10), (80, 50), (99, 99), (29, 29), (50, 50)):
            assert cleaned.getpixel(point) == image.getpixel(point)
    assert raster.data == original
    assert Image.open(io.BytesIO(raster.data)).getpixel((40, 40)) == (0, 0, 0)
    assert read_xml(output).xpath("//w:txbxContent//w:t/text()", namespaces=NS) == ["Исправить"]
    assert any("растр очищен" in warning for warning in warnings)
    assert any("Пересекающиеся линии" in warning for warning in warnings)


def test_ocr_flow_keeps_original_image_for_reference(tmp_path):
    data = io.BytesIO()
    Image.new("RGB", (20, 20), "black").save(data, format="PNG")
    page = PdfPage(1, 300, 400, [TextSpan("Text", (0, 0, 20, 20), "Liberation Sans", 12, (0, 0, 0), from_ocr=True)], [], [RasterImage((0, 0, 20, 20), data.getvalue())], ocr_used=True)
    output = tmp_path / "reference.docx"
    warnings = export_docx([page], output, mode="flow")
    with ZipFile(output) as package:
        media = next(name for name in package.namelist() if name.startswith("word/media/"))
        assert package.read(media) == data.getvalue()
    assert any("для сверки" in warning for warning in warnings)
    assert not any("растр очищен" in warning for warning in warnings)


def test_supplied_font_catalog_retains_user_font_override(tmp_path):
    class CustomCatalog:
        def resolve(self, name, bold=False, italic=False):
            assert name == "Licensed GOST User Font"
            return FontMatch(name, Path("/user/font.ttf"), False, bold, italic)

    output = tmp_path / "custom-font.docx"
    warnings = export_docx([sample_page()], output, font_overrides={"*": "Licensed GOST User Font"}, font_catalog=CustomCatalog())
    assert read_xml(output).xpath("//w:rFonts/@w:hAnsi", namespaces=NS) == ["Licensed GOST User Font"]
    assert not any("заменён" in warning for warning in warnings)


def test_font_override_and_missing_font_warning(tmp_path):
    page = sample_page()
    page.texts[0].font = "NoSuchCommercialGOSTTypeface"
    page.texts[0].bold = False
    output = tmp_path / "fonts.docx"
    warnings = export_docx([page], output)
    assert any("NoSuchCommercialGOSTTypeface" in warning and "заменён" in warning for warning in warnings)
    warnings = export_docx([page], output, font_overrides={"*": "Liberation Serif"})
    assert read_xml(output).xpath("//w:rFonts/@w:hAnsi", namespaces=NS) == ["Liberation Serif"]
    assert not any("NoSuchCommercialGOSTTypeface" in warning for warning in warnings)


def test_large_sheet_scaled_with_text_and_geometry(tmp_path):
    page = sample_page()
    page.width = 2376
    page.height = 3360
    output = tmp_path / "a0.docx"
    warnings = export_docx([page], output)
    assert Document(output).sections[0].page_height.pt == pytest.approx(1584)
    assert any("уменьшена" in warning for warning in warnings)
    style = read_xml(output).xpath("//v:rect/@style", namespaces=NS)[0]
    assert "margin-left:14.1429pt" in style


def test_invalid_xml_controls_are_replaced_and_reported(tmp_path):
    page = sample_page()
    page.texts[0].text = "text\x00remaining"
    output = tmp_path / "control.docx"
    warnings = export_docx([page], output)
    assert read_xml(output).xpath("//w:t/text()", namespaces=NS) == ["text�remaining"]
    assert any("управляющие символы" in warning for warning in warnings)


def test_quarter_turn_text_rotates_about_original_bounding_box_centre(tmp_path):
    page = sample_page()
    page.texts[0].bbox = (20, 40, 35, 140)
    page.texts[0].angle = 90
    output = tmp_path / "rotated.docx"
    export_docx([page], output)
    style = read_xml(output).xpath("//v:rect/@style", namespaces=NS)[0]
    assert "width:100.0000pt;height:15.0000pt" in style
    assert "margin-left:-22.5000pt;margin-top:82.5000pt" in style
    assert "rotation:90.0000" in style


def test_invalid_input_does_not_create_output(tmp_path):
    output = tmp_path / "invalid.docx"
    with pytest.raises(ValueError, match="without pages"):
        export_docx([], output)
    with pytest.raises(ValueError, match="mode"):
        export_docx([sample_page()], output, mode="bad")
    page = sample_page()
    page.width = float("nan")
    with pytest.raises(ValueError, match="page size"):
        export_docx([page], output)
    assert not output.exists()
